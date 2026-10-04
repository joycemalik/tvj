"""
excel_export.py
---------------
Writes outputs/variability_results.xlsx: an organised workbook of the batch run.

Sheets
  README          what each sheet contains, formulas, units, sources
  Fvar by year    one row per (line, year): N, JD range, F avg, F err, F_var ± err
  Light curves    one row per detected measurement: JD, JD − JD_min, F, F err
  Line fits       one row per (spectrum, line): all fitted quantities and verification
  Spectra         one row per spectrum: date, JD, continuum α ± err, A ± err
"""
from typing import Dict, List, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCI = '0.000E+00'
HEADER_FILL = PatternFill('solid', fgColor='E8E6E1')

# (source column, header with units, number format)
FVAR_COLS: List[Tuple[str, str, str]] = [
    ('line', 'Line', '@'), ('year', 'Year', '0'), ('n', 'N spectra', '0'),
    ('jd_min', 'JD min', '0.000'), ('jd_span', 'JD span (d)', '0.0'),
    ('f_mean', 'F avg (erg s-1 cm-2)', SCI), ('err_rms', 'F err rms (erg s-1 cm-2)', SCI),
    ('ratio', 'S²/σ²err', '0.00'), ('fvar', 'Fvar', '0.0000'), ('fvar_err', 'Fvar error', '0.0000'),
    ('regime', 'Error formula used', '@'),
]
CURVE_COLS = [
    ('line', 'Line', '@'), ('year', 'Year', '0'), ('spectrum', 'Spectrum', '@'), ('date', 'Date', '@'),
    ('jd', 'JD', '0.00000'), ('jd_minus_min', 'JD − JD min (d)', '0.000'),
    ('flux', 'F (erg s-1 cm-2)', SCI), ('flux_err', 'F err (erg s-1 cm-2)', SCI),
]
FIT_COLS = [
    ('spectrum', 'Spectrum', '@'), ('date', 'Date', '@'), ('year', 'Year', '0'), ('jd', 'JD', '0.00000'),
    ('line', 'Line', '@'), ('rest_wavelength', 'Rest λ (Å)', '0.00'),
    ('detected', 'Detected (F/σF ≥ 3)', '@'), ('verification', 'Verification', '@'),
    ('failed_checks', 'Failed checks', '@'),
    ('center', 'Line centre μ (Å)', '0.00'), ('center_err', 'μ error (Å)', '0.00'),
    ('sigma', 'σ (Å)', '0.00'), ('sigma_err', 'σ error (Å)', '0.00'),
    ('fwhm_kms', 'FWHM (km/s)', '0'), ('fwhm_kms_intrinsic', 'FWHM intrinsic (km/s)', '0'),
    ('flux', 'Flux (erg s-1 cm-2)', SCI), ('flux_err', 'Flux error', SCI), ('significance', 'F/σF', '0.0'),
    ('flux_window', 'Flux in window (reference convention)', SCI),
    ('ew', 'EW (Å)', '0.00'), ('ew_err', 'EW error (Å)', '0.00'),
    ('wing_window', 'Wing window (Å)', '0'), ('min_wavelength', 'Min λ (Å)', '0.00'),
    ('max_wavelength', 'Max λ (Å)', '0.00'), ('reduced_chi2', 'Reduced χ²', '0.00'),
    ('at_bound', 'Parameter at bound', '@'),
]
SPEC_COLS = [
    ('spectrum', 'Spectrum', '@'), ('date', 'Date', '@'), ('year', 'Year', '0'), ('jd', 'JD', '0.00000'),
    ('spectral_index', 'Spectral index α', '0.0000'), ('spectral_index_err', 'α error', '0.0000'),
    ('continuum_amplitude', 'Amplitude A', SCI), ('continuum_amplitude_err', 'A error', SCI),
    ('n_detected', 'Lines detected', '0'),
]

README = [
    ('3C 273 IUE SWP emission-line variability', ''),
    ('', ''),
    ('Sheet', 'Contents'),
    ('Fvar by year', 'One row per emission line and calendar year. Inputs are the detected (F/σF ≥ 3) line fluxes of that year.'),
    ('Light curves', 'Every detected measurement: JD, JD − JD min (per line, whole campaign), F and F err. Sorted by line wavelength, then JD.'),
    ('Line fits', 'Every spectrum × line: least-squares parameters with 1σ errors and the verification result.'),
    ('Spectra', 'One row per spectrum: observation date and the power-law continuum F = A λ^α.'),
    ('', ''),
    ('Formula', ''),
    ('F avg', 'F̄ = (1/N) Σ F_i'),
    ('S²', 'σ²_F = 1/(N−1) Σ (F_i − F̄)²'),
    ('σ²err', 'ΔF̄² = (1/N) Σ ΔF_i²   (F err rms = √ΔF̄²)'),
    ('Fvar', 'Fvar = √(σ²_F − ΔF̄²) / F̄;  empty when σ²_F ≤ ΔF̄² (scatter consistent with errors)'),
    ('Fvar error (S²/σ²err < 10)', '√(1/2N) · ΔF̄² / (F̄² Fvar)'),
    ('Fvar error (S²/σ²err ≥ 10)', '√(ΔF̄²/N) / F̄'),
    ('Line flux', 'F = √(2π) A σ from a weighted least-squares Gaussian + local linear continuum fit'),
    ('', ''),
    ('Data', 'Rest-frame IUE SWP low-dispersion spectra (z = 0.158); dates from jd.xlsx (SWP log).'),
    ('Sources', 'Vaughan et al. 2003, MNRAS 345, 1271; Sukanya et al. 2018; Rodríguez-Pascual et al. 1997, ApJS 110, 9.'),
    ('Produced by', 'python run_variability.py (see README.md and /about).'),
]


def _sheet(writer, name: str, df: pd.DataFrame, cols) -> None:
    keep = [(c, h, f) for c, h, f in cols if c in df.columns]
    out = df[[c for c, _, _ in keep]].copy()
    out.columns = [h for _, h, _ in keep]
    out.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    ws.freeze_panes = 'A2'
    for j, (c, h, fmt) in enumerate(keep, start=1):
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = max(10, min(42, len(h) + 2, 12 if fmt == SCI else 42))
        if c in ('spectrum', 'failed_checks', 'line', 'regime'):
            ws.column_dimensions[letter].width = 30 if c != 'failed_checks' else 40
        for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
            row[0].number_format = fmt
    ws.row_dimensions[1].height = 32
    ws.auto_filter.ref = ws.dimensions


def write_variability_workbook(path: str, fits: pd.DataFrame, table: pd.DataFrame, curves: pd.DataFrame) -> None:
    order = fits.groupby('line')['rest_wavelength'].first().sort_values().index.tolist()
    rank = {l: i for i, l in enumerate(order)}

    f = fits.copy()
    f['date'] = pd.to_datetime(f['date'], errors='coerce').dt.strftime('%Y-%m-%d')
    f['_r'] = f['line'].map(rank)
    f = f.sort_values(['_r', 'jd', 'spectrum'])
    f['detected'] = f['detected'].map({True: 'yes', False: 'no'})
    f['at_bound'] = f['at_bound'].map({True: 'yes', False: 'no'})

    c = curves.merge(f[['spectrum', 'date']].drop_duplicates('spectrum'), on='spectrum', how='left')
    c['_r'] = c['line'].map(rank)
    c = c.sort_values(['_r', 'jd'])

    t = table.copy()
    t['_r'] = t['line'].map(rank)
    t = t.sort_values(['_r', 'year'])

    s = (fits.assign(date=pd.to_datetime(fits['date'], errors='coerce').dt.strftime('%Y-%m-%d'))
         .groupby('spectrum', as_index=False)
         .agg(date=('date', 'first'), year=('year', 'first'), jd=('jd', 'first'),
              spectral_index=('spectral_index', 'first'), spectral_index_err=('spectral_index_err', 'first'),
              continuum_amplitude=('continuum_amplitude', 'first'),
              continuum_amplitude_err=('continuum_amplitude_err', 'first'),
              n_detected=('detected', 'sum'))
         .sort_values(['jd', 'spectrum'], na_position='last'))

    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        pd.DataFrame(README, columns=['Item', 'Description']).to_excel(xw, sheet_name='README', index=False, header=False)
        ws = xw.sheets['README']
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 110
        for r in (1, 3, 9):
            ws.cell(row=r, column=1).font = Font(bold=True, size=12 if r == 1 else 11)
            ws.cell(row=r, column=2).font = Font(bold=True)
        _sheet(xw, 'Fvar by year', t, FVAR_COLS)
        _sheet(xw, 'Light curves', c, CURVE_COLS)
        _sheet(xw, 'Line fits', f, FIT_COLS)
        _sheet(xw, 'Spectra', s, SPEC_COLS)
